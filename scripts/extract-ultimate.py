#!/usr/bin/env python3
"""Extract the Ultimate DSA Sheet (.xlsx) into scripts/ultimate.json.

One-time data extraction (run with the repo's openpyxl venv):
    .venv/bin/python scripts/extract-ultimate.py "Ultimate DSA Sheet.xlsx"

Header row 18: A=Index, D=Topics, E=Question (hyperlinked), F=Companies, G=Remarks.
Data starts at row 20. The .xlsx itself isn't committed (copyright); the emitted
ultimate.json is, so `npm run migrate-ultimate` works without it.
"""
import json
import sys
from pathlib import Path

import openpyxl

src = sys.argv[1] if len(sys.argv) > 1 else "Ultimate DSA Sheet.xlsx"
wb = openpyxl.load_workbook(src, data_only=True)
ws = wb["Sheet1"]

out = []
n = 0
for r in range(20, ws.max_row + 1):
    topic = ws.cell(r, 4).value
    qcell = ws.cell(r, 5)
    title = qcell.value
    if not title or not str(title).strip():
        continue
    n += 1
    link = qcell.hyperlink.target if qcell.hyperlink else ""
    companies = ws.cell(r, 6).value or ""
    remarks = ws.cell(r, 7).value or ""
    out.append(
        {
            "index": n,
            "topic": str(topic).strip() if topic else "General",
            "title": str(title).strip(),
            "url": (link or "").strip(),
            "companies": str(companies).strip(),
            "remarks": str(remarks).strip(),
        }
    )

dest = Path(__file__).parent / "ultimate.json"
dest.write_text(json.dumps(out, indent=2, ensure_ascii=False))
print(f"wrote {len(out)} problems -> {dest}")
